# -*- coding: utf-8 -*-
"""
정시NAVI 엑셀(경기도교육청 배포) -> 정시 나침반 웹앱용 JSON 변환기.

사용법:  python3 tools/build_jeongsi.py <정시나비.xlsx> [출력디렉터리]
기본 출력: data/jeongsi/

만들어지는 파일
  meta.json    : 등급<->표준점수<->백분위 변환표, 대학별 변환표준점수표, 코드 사전
  units.json   : 모집단위별 전형/반영비율/입시결과 (컬럼 배열 형태로 압축)
"""
import json
import os
import re
import sys
from collections import Counter, defaultdict

import openpyxl

# ---------------------------------------------------------------- 공통 유틸

def s(v):
    """셀 값을 정리된 문자열로."""
    if v is None:
        return ''
    t = str(v).strip()
    if t in ('-', '#N/A', '#VALUE!', '#REF!', '#DIV/0!'):
        return ''
    return re.sub(r'\s+', ' ', t)


def num(v):
    """셀 값을 숫자로. 숫자가 아니면 None."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return round(float(v), 4)
    t = s(v).replace(',', '')
    if not t:
        return None
    try:
        return round(float(t), 4)
    except ValueError:
        return None


def key(univ, term, admit, unit):
    """시트 사이를 잇는 조인 키.

    전형명 표기가 시트마다 흔들린다. '일반학생전형'과 '일반학생', '일반전형1'과 '일반1'을
    같은 것으로 보게 끝(또는 끝의 번호 앞)에 붙은 '전형'을 떼어 낸다.
    """
    def n(x):
        x = s(x).replace(' ', '')
        return re.sub(r'전형(?=\d*$)', '', x)
    return (n(univ), n(term), n(admit), n(unit))


# ------------------------------------------------- 1. 등급/표준점수/백분위 변환표

def build_score_tables(wb):
    """'25변환' 시트에서 과목별 (표준점수, 백분위, 등급) 분포를 읽는다.

    2026 수능 실제 분포이므로 목표 등급 -> 대표 백분위/표준점수 환산의 근거가 된다.
    """
    ws = wb['25변환']
    raw = defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        name = s(r[0])
        if not name:
            continue
        m = re.match(r'^(국어|수학|과탐|사과탐)(\d+)$', name)
        if not m:
            continue
        std, pct, grade = num(r[1]), num(r[2]), num(r[3])
        if std is None or pct is None or grade is None:
            continue
        raw[m.group(1)].append([int(std), int(pct), int(grade)])

    tables = {}
    for sub, rows in raw.items():
        rows.sort(key=lambda x: -x[0])          # 표준점수 내림차순
        tables[sub] = rows
    return tables



def pct_to_std(tables):
    """백분위(0~100) -> 대표 표준점수. 대학이 표준점수를 활용할 때 쓴다."""
    out = {}
    for sub, rows in tables.items():
        by = defaultdict(list)
        for std, pct, g in rows:
            by[pct].append(std)
        arr = []
        for p in range(101):
            if p in by:
                v = sorted(by[p])
                arr.append(v[len(v) // 2])
            else:
                arr.append(None)
        # 빈 구간은 이웃 값으로 선형 보간
        known = [i for i, v in enumerate(arr) if v is not None]
        for i, v in enumerate(arr):
            if v is not None:
                continue
            lo = max([k for k in known if k < i], default=None)
            hi = min([k for k in known if k > i], default=None)
            if lo is None:
                arr[i] = arr[hi]
            elif hi is None:
                arr[i] = arr[lo]
            else:
                t = (i - lo) / (hi - lo)
                arr[i] = round(arr[lo] + (arr[hi] - arr[lo]) * t, 2)
        out[sub] = arr
    return out


def grade_bands(tables):
    """등급별 (표준점수, 백분위)의 상단/중앙/하단 대표값."""
    out = {}
    for sub, rows in tables.items():
        by_grade = defaultdict(list)
        for std, pct, g in rows:
            by_grade[g].append((std, pct))
        band = {}
        for g, vals in by_grade.items():
            vals.sort(key=lambda x: -x[0])
            top, mid, low = vals[0], vals[len(vals) // 2], vals[-1]
            band[g] = {'top': list(top), 'mid': list(mid), 'low': list(low)}
        out[sub] = band
    return out


def build_byunpyo(wb):
    """'변표' 시트에서 대학별 2026 변환표준점수표(백분위 0~100)를 읽는다."""
    ws = wb['변표']
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    tables = {}
    i = 0
    while i < len(rows):
        head = rows[i]
        # 헤더행: C열에 '통합' / '사탐' 등 구분 문자열이 있고, 다음 행부터 데이터
        if isinstance(head[2], str) and s(head[2]) and s(head[2]) != '점수':
            nxt = rows[i + 1] if i + 1 < len(rows) else None
            name = ''
            if nxt and isinstance(nxt[1], str):
                m = re.match(r'^(.*?)(\d+)$', s(nxt[1]))
                name = m.group(1) if m else ''
            if name:
                cols = [s(head[c]) for c in (2, 3, 4, 5)]
                series = {c: [None] * 101 for c in cols if c}
                j = i + 1
                while j < len(rows):
                    rr = rows[j]
                    p = num(rr[0])
                    if p is None or not isinstance(rr[1], str) or not s(rr[1]).startswith(name):
                        break
                    for ci, cname in zip((2, 3, 4, 5), cols):
                        if cname:
                            v = num(rr[ci])
                            if v is not None and 0 <= int(p) <= 100:
                                series[cname][int(p)] = v
                    j += 1
                tables[name] = {k: v for k, v in series.items() if any(x is not None for x in v)}
                i = j
                continue
        i += 1
    return tables


# ------------------------------------------------------------ 2. 보조 시트 읽기

def read_grade_points(wb, sheet):
    """'영어'/'한국사' 시트: 모집단위별 등급(1~9) 반영점수와 반영방식."""
    ws = wb[sheet]
    out = {}
    for r in ws.iter_rows(min_row=12, values_only=True):
        if not s(r[4]):
            continue
        k = key(r[4], r[5], r[7], r[8])
        pts = [num(r[c]) for c in range(14, 23)]
        if all(p is None for p in pts):
            pts = None
        out[k] = {
            'method': s(r[11]),
            'ratio': s(r[12]),
            'pts': pts,
            'impact': [num(r[23]), num(r[24]), num(r[25])],
        }
    return out


def read_ratio_sheet(wb, sheet, year=2026):
    """'반영비율(모집단위)' / '반영비율(모집계열)' 시트 -> 조인키별 반영 정보."""
    ws = wb[sheet]
    out = {}
    for r in ws.iter_rows(min_row=8, values_only=True):
        if not s(r[3]) or num(r[1]) != year:
            continue
        rec = {
            'region': s(r[2]),
            'univ': s(r[3]),
            'term': s(r[4]),
            'admit': s(r[5]),
            'unit': s(r[6]),
            'idxKM': s(r[7]),          # 국어·수학 활용지표
            'idxT': s(r[8]),           # 탐구 활용지표
            'totalAll': num(r[9]),     # 전형 총점
            'total': num(r[10]),       # 수능 총점
            'areas': s(r[11]),         # 반영영역 표기
            'areaCnt': num(r[12]),
            'ratioChanged': s(r[13]),
            'w': [s(r[14]), s(r[15]), s(r[16]), s(r[17]), s(r[19]), s(r[20])],  # 국 수 영 탐 한 제2
            'tCnt': num(r[18]),
            'reqMath': s(r[21]),
            'reqTam': s(r[22]),
            'engMethod': s(r[23]),
            'hisMethod': s(r[24]),
            'bonusArea': s(r[25]),
            'bonus': s(r[26]),
            'note': s(r[27]),
        }
        out[key(r[3], r[4], r[5], r[6])] = rec
    return out


def read_compete(wb):
    """'경쟁률|충원율' 시트: 최근 3개년 경쟁률·충원율·실질경쟁률.

    실질경쟁률 = 경쟁률 / (1 + 충원율). 추가합격이 도는 만큼 실제 문턱은 낮아진다.
    """
    ws = wb['경쟁률|충원율 ']
    out = {}
    for r in ws.iter_rows(min_row=9, values_only=True):
        if not s(r[10]):
            continue
        k = key(r[10], '', r[11], r[12])
        out[k] = [num(r[14]), num(r[15]), num(r[16]),      # 2025 경쟁률·충원율·실질
                  num(r[18]), num(r[19]), num(r[20]),      # 2024
                  num(r[22]), num(r[23]), num(r[24])]      # 2023
    return out


def read_move(wb):
    """'군이동' 시트: 대학별 가·나·다군 모집정원 증감.

    군이 바뀌면 지원자 풀이 통째로 달라져 입시결과가 크게 흔들린다.
    """
    ws = wb['군이동']
    out = {}
    for r in ws.iter_rows(min_row=9, values_only=True):
        name = s(r[2])
        if not name or name == '*':
            continue
        out[name] = {
            'pct': num(r[3]),                                    # 대학 전체 70%컷 백분위 평균
            'seats': num(r[4]), 'seatDiff': num(r[5]),
            'y26': [num(r[6]), num(r[7]), num(r[8])],            # 2026 가·나·다
            'y25': [num(r[9]), num(r[10]), num(r[11])],
            'diff': [num(r[12]), num(r[13]), num(r[14])],
            'termPct': [num(r[18]), num(r[19]), num(r[20])],     # 군별 백분위 평균
        }
    return out


# ------------------------------------------------------------------ 3. 본 처리

REGION_ORDER = ['서울', '경인권', '충청권', '강원권', '전라권', '경상권', '제주권']


def build(xlsx_path, out_dir):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    score_tables = build_score_tables(wb)
    bands = grade_bands(score_tables)
    byunpyo = build_byunpyo(wb)
    eng = read_grade_points(wb, '영어')
    his = read_grade_points(wb, '한국사')
    ratio_unit = read_ratio_sheet(wb, '반영비율(모집단위)')
    ratio_track = read_ratio_sheet(wb, '반영비율(모집계열)')
    compete = read_compete(wb)
    move = read_move(wb)

    # 폴백용 인덱스: (대학, 군, 모집단위) / (대학, 군, 전형) / (대학, 군)
    idx_unit = defaultdict(list)
    idx_admit = defaultdict(list)
    for k, v in ratio_unit.items():
        idx_unit[(k[0], k[1], k[3])].append(v)
        idx_admit[(k[0], k[1], k[2])].append(v)
    idx_track = defaultdict(list)
    for k, v in ratio_track.items():
        idx_track[(k[0], k[1], k[2])].append(v)
        idx_track[(k[0], k[1])].append(v)

    def same(cands):
        return len({(x['areas'], tuple(x['w']), x['total']) for x in cands}) == 1

    def narrow(cands, unit_name, track):
        """모집계열 시트 행 가운데 이 모집단위/계열에 해당하는 것만 남긴다."""
        if not cands:
            return cands
        hit = [x for x in cands if unit_name and unit_name in (x['unit'] or '')]
        if hit:
            return hit
        if track:
            hit = [x for x in cands if track and track in (x['unit'] or '')]
            if hit:
                return hit
        hit = [x for x in cands if '전체' in (x['unit'] or '')]
        return hit or cands

    def lookup_ratio(k, unit_name='', track=''):
        if k in ratio_unit:
            return ratio_unit[k], 'exact'
        c = idx_unit.get((k[0], k[1], k[3]))
        if c and same(c):
            return c[0], 'unit'
        c = idx_admit.get((k[0], k[1], k[2]))
        if c and same(c):
            return c[0], 'admit'
        c = narrow(idx_track.get((k[0], k[1], k[2])) or idx_track.get((k[0], k[1])), unit_name, track)
        if c and same(c):
            return c[0], 'track'
        if c:
            return c[0], 'track~'
        c = idx_admit.get((k[0], k[1], k[2]))
        if c:
            return c[0], 'admit~'
        return None, None

    def lookup_pts(table, k):
        if k in table:
            return table[k]
        return None

    def req_of(k, rt, how, field):
        """지정과목은 잘못 물려받으면 지원 가능한 곳을 통째로 지워 버린다.
        정확히 맞은 행이 아니면, 같은 대학·전형 안에서 값이 하나로 모일 때만 쓴다."""
        if not rt:
            return ''
        if how == 'exact':
            return rt.get(field, '')
        pool = idx_admit.get((k[0], k[1], k[2])) or []
        vals = {x.get(field, '') for x in pool if x.get(field, '')}
        if len(vals) == 1:
            return vals.pop()
        if len(vals) > 1:
            return ''          # 계열마다 달라 단정할 수 없다 -> 제한 없음으로 둔다
        return rt.get(field, '')

    ws = wb['2026정시']
    units = []
    stat = Counter()
    for r in ws.iter_rows(min_row=7, values_only=True):
        if not s(r[3]):
            continue
        k = key(r[3], r[4], r[5], r[6])
        rt, how = lookup_ratio(k, s(r[6]), s(r[11]))
        stat[how or 'none'] += 1
        e = lookup_pts(eng, k) or {}
        h = lookup_pts(his, k) or {}
        cp = compete.get(key(r[3], '', r[5], r[6])) or [None] * 9

        units.append({
            'zone': s(r[0]), 'region': s(r[1]), 'city': s(r[2]),
            'univ': s(r[3]), 'term': s(r[4]), 'admit': s(r[5]),
            'unit': s(r[6]), 'major': s(r[7]),
            'chgAdmit': s(r[8]), 'chgUnit': s(r[9]), 'freeMajor': s(r[10]),
            'track': s(r[11]),
            'n26': num(r[12]),
            'method': s(r[13]),
            'idxKM': s(r[14]) or (rt or {}).get('idxKM', ''),
            'idxT': s(r[15]) or (rt or {}).get('idxT', ''),
            'tCnt': num(r[16]) or (rt or {}).get('tCnt'),
            'reqTamRaw': s(r[17]),
            'memo': s(r[18]),
            'wRaw': [s(r[19]), s(r[20]), s(r[21]), s(r[22]), s(r[24]), s(r[23])],
            'n25': num(r[29]), 'carry25': num(r[30]), 'nf25': num(r[31]),
            'rate25': num(r[32]), 'fill25': num(r[33]),
            'cut50': num(r[34]), 'cut70': num(r[35]),
            'total': num(r[36]) or (rt or {}).get('total'),
            'totalAll': num(r[37]) or (rt or {}).get('totalAll'),
            'cutPct': num(r[53]),
            'ratioChanged': s(r[54]),
            'nf24': num(r[55]), 'rate24': num(r[56]), 'fill24': num(r[57]),
            # 반영비율 시트에서 온 정밀 정보
            'areas': (rt or {}).get('areas', ''),
            'w': (rt or {}).get('w') or [s(r[19]), s(r[20]), s(r[21]), s(r[22]), s(r[24]), s(r[23])],
            'reqMath': req_of(k, rt, how, 'reqMath'),
            'reqTam': req_of(k, rt, how, 'reqTam') or s(r[17]),
            'bonus': (rt or {}).get('bonus', ''),
            'bonusArea': (rt or {}).get('bonusArea', ''),
            'note': (rt or {}).get('note', ''),
            'engMethod': e.get('method') or (rt or {}).get('engMethod', ''),
            'engPts': e.get('pts'),
            'engImpact': e.get('impact'),
            'hisMethod': h.get('method') or (rt or {}).get('hisMethod', ''),
            'hisPts': h.get('pts'),
            'match': how or '',
            'compete': cp,
            'move': move.get(s(r[3])),
        })

    print('반영비율 매칭:', stat.most_common())
    print('모집단위 수:', len(units))

    os.makedirs(out_dir, exist_ok=True)
    meta = {
        'source': os.path.basename(xlsx_path),
        'scoreTables': score_tables,
        'pctToStd': pct_to_std(score_tables),
        'maxStd': {k: v[0][0] for k, v in score_tables.items()},
        'gradeBands': bands,
        'byunpyo': byunpyo,
        'regionOrder': REGION_ORDER,
    }
    with open(os.path.join(out_dir, 'meta.json'), 'w', encoding='utf-8') as f:
        json.dump(meta, f, ensure_ascii=False, separators=(',', ':'))
    with open(os.path.join(out_dir, 'units.json'), 'w', encoding='utf-8') as f:
        json.dump(compact(units), f, ensure_ascii=False, separators=(',', ':'))
    if os.environ.get('KEEP_RAW'):
        with open(os.path.join(out_dir, 'units.raw.json'), 'w', encoding='utf-8') as f:
            json.dump(units, f, ensure_ascii=False, separators=(',', ':'))
    print('저장 완료 ->', out_dir)
    return units, meta


# --------------------------------------------------------------- 4. 압축 저장

COLS = ['zone', 'region', 'city', 'univ', 'term', 'admit', 'unit', 'major', 'track',
        'n26', 'idxKM', 'idxT', 'tCnt', 'areas', 'w', 'reqMath', 'reqTam',
        'engMethod', 'engPts', 'engImpact', 'hisMethod', 'hisPts',
        'total', 'totalAll', 'cut50', 'cut70', 'cutPct',
        'n25', 'carry25', 'nf25', 'rate25', 'fill25', 'rate24', 'fill24',
        'chgAdmit', 'chgUnit', 'ratioChanged', 'freeMajor', 'method',
        'memo', 'note', 'bonus', 'bonusArea', 'match', 'compete', 'move']

DICT_COLS = {'zone', 'region', 'city', 'univ', 'term', 'admit', 'unit', 'major', 'track',
             'idxKM', 'idxT', 'areas', 'w', 'reqMath', 'reqTam', 'engMethod', 'engPts',
             'engImpact', 'hisMethod', 'hisPts', 'method', 'memo', 'note', 'bonus',
             'bonusArea', 'chgAdmit', 'chgUnit', 'ratioChanged', 'freeMajor', 'match',
             'compete', 'move'}


def compact(units):
    """반복이 심한 열을 사전(인덱스) 방식으로 압축한다."""
    dicts = {c: {} for c in DICT_COLS}
    order = {c: [] for c in DICT_COLS}

    def idx(col, val):
        k = json.dumps(val, ensure_ascii=False, separators=(',', ':'))
        d = dicts[col]
        if k not in d:
            d[k] = len(order[col])
            order[col].append(val)
        return d[k]

    rows = []
    for u in units:
        row = []
        for c in COLS:
            v = u.get(c)
            if c in DICT_COLS:
                row.append(idx(c, v if v is not None else ''))
            else:
                row.append(v)
        rows.append(row)
    return {'cols': COLS, 'dict': order, 'rows': rows}


if __name__ == '__main__':
    src = sys.argv[1] if len(sys.argv) > 1 else 'navi.xlsx'
    dst = sys.argv[2] if len(sys.argv) > 2 else 'data/jeongsi'
    build(src, dst)
